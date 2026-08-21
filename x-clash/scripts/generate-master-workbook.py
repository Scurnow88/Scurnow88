#!/usr/bin/env python3
"""Build X-CLASH-MASTER-RAW-DATA.xlsx from CSV sources."""

from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]

SHEETS = [
    ("My Teams", "x-clash-my-teams.csv"),
    ("My Roster", "x-clash-my-roster.csv"),
    ("Team Compositions", "x-clash-team-compositions.csv"),
    ("Hero Slot Guide", "x-clash-hero-slot-guide.csv"),
    ("Upgrade Priority", "x-clash-upgrade-priority.csv"),
    ("Faction Reference", "x-clash-faction-reference.csv"),
    ("Skills Spec", "x-clash-heroes-skills-spec.csv"),
    ("Heroes Master", "x-clash-heroes-master.csv"),
    ("Tier By Mode", "x-clash-heroes-tier-by-mode.csv"),
    ("Skills Raw", "x-clash-heroes-skills.csv"),
    ("Roster Template", "x-clash-my-roster-template.csv"),
]


def load_csv(path: Path) -> list[list[str]]:
    rows: list[list[str]] = []
    text = path.read_text(encoding="utf-8")
    for line in text.splitlines():
        if not line.strip():
            continue
        rows.append([cell.strip().strip('"') for cell in _parse_csv_line(line)])
    return rows


def _parse_csv_line(line: str) -> list[str]:
    """Minimal CSV parser (no quoted commas in our data)."""
    out: list[str] = []
    cur: list[str] = []
    in_quotes = False
    for ch in line:
        if ch == '"':
            in_quotes = not in_quotes
        elif ch == "," and not in_quotes:
            out.append("".join(cur))
            cur = []
        else:
            cur.append(ch)
    out.append("".join(cur))
    return out


def autosize_columns(ws) -> None:
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for row in range(1, min(ws.max_row, 200) + 1):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[letter].width = min(max_len + 2, 48)


def main() -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for title, filename in SHEETS:
        path = ROOT / filename
        if not path.exists():
            print(f"skip missing: {filename}")
            continue
        ws = wb.create_sheet(title)
        for r_idx, row in enumerate(load_csv(path), start=1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
        autosize_columns(ws)
        print(f"  + {title} ({path.name})")

    out = ROOT / "X-CLASH-MASTER-RAW-DATA.xlsx"
    wb.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
